"""
Build the material decarbonisation paths of Horup et al. (2025) from their
supplementary material. No values typed by hand.

Source (provenance in data/raw/MANIFEST.csv; the file itself is kept in refs/,
which is git-ignored):
  refs/Horup_2025_SI_mmc1.xlsx
  Horup L, Bruhn S, Hoxha E, Birgisdottir H, Secher AQ, Ohms P, Hauschild M
  (2025). Absolute sustainability assessment of the Danish building sector
  through prospective LCA. Sci. Total Environ. 966, 178780. Supplementary
  material (mmc1.xlsx), sheet 'Table S4' (in-sheet title 'Table S3: Material
  analysis (extended) of GWP by applying prospective background system').

For each material the sheet gives, in column C, the premise result (average of
the 9 scenarios) relative to ecoinvent 3.9.1 in 2025, 2030, 2040 and 2050, and
in column E the 'average mitigation potential 2025-2050'. The path relative to
2025 is C(t) / C(2025); the script checks that 1 - C(2050)/C(2025) equals
column E for every material.

Output: data/derived/horup2025_table_s4.csv
  material, year, premise_rel_ecoinvent, rel_to_2025, cell
"""
import os

import openpyxl
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(os.path.dirname(HERE), 'refs', 'Horup_2025_SI_mmc1.xlsx')
OUT = os.path.join(HERE, 'derived', 'horup2025_table_s4.csv')
SHEET = 'Table S4'
YEARS = (2025, 2030, 2040, 2050)


def main():
    if not os.path.exists(SRC):
        raise SystemExit(f'{SRC} not found (refs/ is not versioned; see data/raw/MANIFEST.csv).')
    ws = openpyxl.load_workbook(SRC, data_only=True)[SHEET]
    if 'Premise (average) relative to Ecoinvent' not in str(ws['C2'].value):
        raise ValueError(f'Unexpected header C2 = {ws["C2"].value!r}.')
    rows, cur, mitigation = [], None, {}
    for r in range(1, ws.max_row + 1):
        name, year, val = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if year == YEARS[0]:
            cur = str(name).strip()
        if cur and year in YEARS:
            rows.append(dict(material=cur, year=int(year), premise_rel_ecoinvent=float(val), cell=f'C{r}'))
            if year == YEARS[-1]:
                mitigation[cur] = float(ws.cell(r, 5).value)
    d = pd.DataFrame(rows)
    if (d.groupby('material').size() != len(YEARS)).any():
        raise ValueError('A material lacks one of the years 2025/2030/2040/2050.')
    base = d[d['year'] == YEARS[0]].set_index('material')['premise_rel_ecoinvent']
    d['rel_to_2025'] = d['premise_rel_ecoinvent'] / d['material'].map(base)
    r50 = d[d['year'] == YEARS[-1]].set_index('material')['rel_to_2025']
    worst = max(abs((1 - r50[m]) - mitigation[m]) for m in r50.index)
    if worst > 1e-12:
        raise ValueError(f'1 - C(2050)/C(2025) differs from column E by {worst:g}.')
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    d.to_csv(OUT, index=False)
    print(f'{os.path.basename(SRC)} [{SHEET}] -> {OUT}: {d.material.nunique()} materials; '
          f'1 - C2050/C2025 = column E for all (max difference {worst:.1e})')


if __name__ == '__main__':
    main()
