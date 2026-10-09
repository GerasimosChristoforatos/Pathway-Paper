"""Planetary-boundary layer (boundaries.py, aesa.py): equations, reproduction of the
earlier manuscript and of its constants, identities of the method, units.
Tests that read outputs/pb_*.csv / pb_results.json rely on run_all.py having
run boundaries.py before the tests (as it does)."""
import json
import os

import numpy as np
import openpyxl
import pandas as pd
import pytest

from conftest import ROOT

import aesa  # noqa: E402
import boundaries  # noqa: E402

WORKBOOK = os.path.join(ROOT, 'data', 'raw', 'Absolute_Analysis.xlsm')   # local only (git-ignored)
OUT = os.path.join(ROOT, 'outputs')


def results():
    with open(os.path.join(OUT, 'pb_results.json')) as f:
        return json.load(f)


def annual():
    return pd.read_csv(os.path.join(OUT, 'pb_annual.csv'))


# ---- reproduction of the earlier manuscript ----------------------------------
@pytest.mark.skipif(not os.path.exists(WORKBOOK), reason='Absolute_Analysis.xlsm not present (kept local only)')
def test_manuscript_ar_boundaries_2025_2029():
    """Sheet '1.Acquired Rights', GWP block (rows 10-14): G = F*(E/D)*((C*$B$10)/(B*$C$10)),
    reproduced from the sheet's own inputs with aesa.boundary_ar."""
    ws_f = openpyxl.load_workbook(WORKBOOK, data_only=False)['1.Acquired Rights']
    ws_v = openpyxl.load_workbook(WORKBOOK, data_only=True)['1.Acquired Rights']
    assert ws_f['G10'].value.replace(' ', '') == '=F10*(E10/D10)*((C10*$B$10)/(B10*$C$10))'
    p_w0, p_nz0 = ws_v['B10'].value, ws_v['C10'].value
    for r in range(10, 15):
        v = {c: ws_v[f'{c}{r}'].value for c in 'ABCDEFG'}
        got = aesa.boundary_ar(v['F'], v['E'], v['D'], v['C'], v['B'], p_nz0, p_w0)
        assert got == pytest.approx(v['G'], rel=1e-12), (v['A'], got, v['G'])


@pytest.mark.skipif(not os.path.exists(WORKBOOK), reason='Absolute_Analysis.xlsm not present (kept local only)')
def test_manuscript_ssec_constants_match_workbook():
    ws = openpyxl.load_workbook(WORKBOOK, data_only=True)['2.Equalitarian']
    ss = boundaries.sector_share_components()
    assert ss['consented_value_nzd_per_yr'] == pytest.approx(ws['F2'].value, rel=1e-12)
    assert ss['gdp_nzd_per_yr'] == pytest.approx(ws['F3'].value, rel=1e-12)
    assert ss['fce_share_mean_of_ratios'] == pytest.approx(ws['F4'].value, rel=1e-12)


def test_ssec_from_raw_files():
    """s_sec from data/raw/GDP_Data.xlsx (via data/derived) and the consents:
    the manuscript's constants (author's brief) and s_sec = 0.0525462."""
    ss = boundaries.sector_share_components()
    assert ss['consented_value_nzd_per_yr'] == pytest.approx(1.42898509115e10, rel=1e-12)
    assert ss['gdp_nzd_per_yr'] == pytest.approx(3.523043e11, rel=1e-12)
    assert ss['fce_share_mean_of_ratios'] == pytest.approx(0.771913492, abs=1e-9)
    assert ss['fce_share_ratio_of_sums'] == pytest.approx(0.773709, abs=5e-7)
    assert ss['s_sec_mean_of_ratios'] == pytest.approx(0.0525462, abs=5e-8)


# ---- equations ------------------------------------------------------------------
def test_alpha_equals_eq_s1_with_totals():
    rng = np.random.default_rng(1)
    p = rng.uniform(1e5, 1e9, 50)
    g = rng.uniform(500, 1e5, 50)
    assert aesa.alpha(p, g) == pytest.approx(aesa.alpha_eq_s1(p, p * g), rel=1e-12)


def test_atp_normalisation_sums_to_sos():
    """Sum over covered countries of SoSOS_ind,c x P_c = SOS (Hjalsted SI Eq. S1),
    with P_W and g_W from the same covered set."""
    r = results()
    b = r['components']['base_year_b']
    w = boundaries.wdi_table()
    d = w[(w['year'] == b) & w['SP.POP.TOTL'].notna() & w[boundaries.GDP_MEASURE].notna()]
    p, g = d['SP.POP.TOTL'].values, d[boundaries.GDP_MEASURE].values
    a, gw = aesa.alpha(p, g), aesa.world_gdp_per_capita(p, g)
    sos = boundaries.SOURCE_PAIRS[boundaries.SOURCE_PAIR]['sos']
    total = sum(aesa.sosos_ind_atp(sos, p.sum(), gw, gc, a) * pc for pc, gc in zip(p, g))
    assert total == pytest.approx(sos, rel=1e-12)
    assert a == pytest.approx(r['components']['alpha'], rel=1e-12)


def test_alpha_2017_reported_against_hjalsted():
    """alpha(2017) is reported next to Hjalsted's 0.243; it is not tuned to it."""
    c = results()['alpha_check']
    assert c['hjalsted_alpha'] == 0.243 and c['year'] == 2017
    for m in (boundaries.GDP_MEASURE, boundaries.GDP_MEASURE_SENS):
        assert 0 < c[m]['alpha'] < 1


def test_population_modifier_is_one_at_t0():
    assert aesa.population_modifier(5.3e6, 8.2e9, 5.3e6, 8.2e9) == 1.0
    a = annual()
    t0 = a[(a['year'] == boundaries.T0) & (a['scenario'] != 'history')]
    assert np.allclose(t0['pop_modifier'], 1.0, rtol=0, atol=1e-15)


def test_ar_identity_over_base_window(B):
    """LoE_AR over the base window = I_glob / SOS by construction: the mean
    base-window impact over B_AR(t0) (modifier = 1) cancels the sector impact."""
    hist = boundaries.historical_emissions(B, boundaries.AR_BASE_WINDOW_RANGE())
    i_base = aesa.mean_over(hist['I_sec_kg'], boundaries.AR_BASE_WINDOW)
    pr = boundaries.SOURCE_PAIRS[boundaries.SOURCE_PAIR]
    b_t0 = aesa.boundary_ar(pr['sos'], i_base, pr['i_glob'], 1.0, 1.0, 1.0, 1.0)
    assert aesa.cumulative_loe(hist['I_sec_kg'], np.full(len(hist), b_t0)) == pytest.approx(
        pr['i_glob'] / pr['sos'], rel=1e-12)
    assert i_base == pytest.approx(results()['components']['I_sec_base_kg'], rel=1e-12)


def test_history_decomposition(B):
    """Soil on all floor area reproduces the diag_5 history (2025 factors x built
    floor area); the projection's convention removes soil only on net replacement."""
    h = boundaries.historical_emissions(B, range(2016, 2026))
    base = B['UPFRONT_2025'] if B['TIME_RESOLVED'] else B['T_BASELINE_2025']
    direct = sum(B['hist_built_typ_gfa'].loc[2016:2025, t] * base[t] for t in B['typ_names']) + h['b_kg']
    assert np.allclose(h['I_soil_on_all_kg'].values, direct.values, rtol=1e-12)
    assert ((h['soil_share'] > 0) & (h['soil_share'] <= 1)).all()
    assert (h['I_sec_kg'] <= h['I_soil_on_all_kg'] + 1e-6).all()


def test_atp_over_ar_constant():
    a = annual()
    r = a[a['scenario'] == 'REF']
    q = r['B_AtP_kgCO2e'] / r['B_AR_kgCO2e']
    assert q.std() / q.mean() < 1e-12


def test_loe_definitions():
    a = annual()
    r = a[(a['scenario'] == 'REF') & (a['year'] >= boundaries.PB_YEARS[0])]
    assert np.allclose(r['LoE_AR'], r['I_sec_kgCO2e'] / r['B_AR_kgCO2e'], rtol=1e-12)
    cum = r['I_sec_kgCO2e'].sum() / r['B_AR_kgCO2e'].sum()
    assert cum == pytest.approx(results()['deterministic']['REF']['LoE_AR'], rel=1e-12)


def test_units_one_year_by_hand():
    """B_AR(2030) from its components in kg, against pb_annual (kg) and pb_results (kt)."""
    r, a = results(), annual()
    c = r['components']
    row = a[(a['scenario'] == 'REF') & (a['year'] == 2030)].iloc[0]
    by_hand_kg = c['SOS_kg'] * c['I_sec_base_kg'] / c['I_glob_kg'] * (row['P_NZ'] / c['P_NZ_t0']) \
        * (c['P_W_t0'] / row['P_W'])
    assert row['B_AR_kgCO2e'] == pytest.approx(by_hand_kg, rel=1e-12)
    assert r['boundaries']['2030']['B_AR_kt'] * 1e6 == pytest.approx(by_hand_kg, rel=1e-12)
    atp_kg = c['SOS_kg'] / row['P_W'] / c['g_NZ'] * c['g_W'] * c['alpha'] * row['P_NZ'] * c['s_sec']
    assert row['B_AtP_kgCO2e'] == pytest.approx(atp_kg, rel=1e-12)


def test_mc_percentiles_ordered():
    m = pd.read_csv(os.path.join(OUT, 'pb_mc_annual.csv'))
    for k in ('LoE_AR', 'LoE_AtP', 'B_AR_kgCO2e', 'I_sec_kgCO2e'):
        assert (m[f'{k}_p5'] <= m[f'{k}_p50']).all() and (m[f'{k}_p50'] <= m[f'{k}_p95']).all()
    assert set(m['scenario']) == {'REF', 'LOW', 'HIGH'}
